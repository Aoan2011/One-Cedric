"""Word 文档操作：python-docx + win32com。"""
from __future__ import annotations

from pathlib import Path

from ..sandbox import _resolve_path
from .common import (
    IS_WINDOWS, check_file, format_size, system_print,
    win32_com_available, encrypt_with_msoffcrypto,
)

DOCX_EXTS = (".docx", ".doc")


def _docx_read(root: Path, args: dict) -> str:
    try:
        from docx import Document
    except ImportError:
        return "ERROR: 需要 python-docx（pip install python-docx）"

    path = str(args.get("path", ""))
    p, err = _resolve_path(root, path)
    if err:
        return err
    err = check_file(p, DOCX_EXTS)
    if err:
        return err

    try:
        doc = Document(str(p))
    except Exception as exc:
        return f"ERROR: 打开文档失败: {exc}"

    lines = [f"文件: {p.name}", f"大小: {format_size(p.stat().st_size)}", ""]

    paras = [para.text for para in doc.paragraphs if para.text.strip()]
    lines.append(f"段落数: {len(paras)}")
    lines.append(f"表格数: {len(doc.tables)}")
    lines.append(f"内联图片数: {len(doc.inline_shapes)}")
    lines.append("")
    lines.append("--- 正文（前 100 段） ---")
    for i, t in enumerate(paras[:100], 1):
        lines.append(f"{i:>3}| {t[:200]}")
    if len(paras) > 100:
        lines.append(f"... 还有 {len(paras) - 100} 段")

    if doc.tables:
        lines.append("")
        lines.append("--- 表格摘要 ---")
        for ti, table in enumerate(doc.tables[:10], 1):
            rows = len(table.rows)
            cols = len(table.columns) if rows else 0
            lines.append(f"表 {ti}: {rows} 行 × {cols} 列")
            for ri, row in enumerate(table.rows[:3], 1):
                cells = [c.text.strip()[:20] for c in row.cells]
                lines.append(f"  {ri}| " + " | ".join(cells))

    return "\n".join(lines)


def _docx_extract_images(root: Path, args: dict) -> str:
    try:
        from docx import Document
    except ImportError:
        return "ERROR: 需要 python-docx"

    path = str(args.get("path", ""))
    out_dir = str(args.get("out_dir", "") or "")

    p, err = _resolve_path(root, path)
    if err:
        return err
    err = check_file(p, DOCX_EXTS)
    if err:
        return err

    if not out_dir:
        out_dir = f"{p.stem}_images"
    od, err = _resolve_path(root, out_dir)
    if err:
        return err

    try:
        doc = Document(str(p))
    except Exception as exc:
        return f"ERROR: 打开失败: {exc}"

    od.mkdir(parents=True, exist_ok=True)
    count = 0
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            try:
                img_data = rel.target_part.blob
                ext = rel.target_part.partname.suffix or ".png"
                fname = f"image_{count + 1}{ext}"
                (od / fname).write_bytes(img_data)
                count += 1
            except Exception:
                continue

    if count == 0:
        return f"未找到图片: {p.name}"
    return f"已提取 {count} 张图片到 {od.relative_to(root)}"


def _docx_edit(root: Path, args: dict) -> tuple[str, bool, str]:
    try:
        from docx import Document
    except ImportError:
        return "", False, "ERROR: 需要 python-docx"

    path = str(args.get("path", ""))
    find_text = str(args.get("find", "") or "")
    replace_text = str(args.get("replace", "") or "")
    append_text = str(args.get("append", "") or "")

    if not find_text and not append_text:
        return "", False, "ERROR: 需要 find/replace 或 append 参数"

    p, err = _resolve_path(root, path)
    if err:
        return "", False, err
    err = check_file(p, DOCX_EXTS)
    if err:
        return "", False, err

    preview_parts = [f"将修改: {p.name}"]
    if find_text:
        preview_parts.append(f"  替换 '{find_text[:50]}' → '{replace_text[:50]}'")
    if append_text:
        preview_parts.append(f"  追加 {len(append_text)} 字符")
    return "\n".join(preview_parts), True, ""


def _docx_edit_apply(root: Path, args: dict) -> str:
    try:
        from docx import Document
    except ImportError:
        return "ERROR: 需要 python-docx"

    path = str(args.get("path", ""))
    find_text = str(args.get("find", "") or "")
    replace_text = str(args.get("replace", "") or "")
    append_text = str(args.get("append", "") or "")

    p, err = _resolve_path(root, path)
    if err:
        return err

    try:
        doc = Document(str(p))
    except Exception as exc:
        return f"ERROR: 打开失败: {exc}"

    replaced = 0
    if find_text:
        for para in doc.paragraphs:
            if find_text in para.text:
                for run in para.runs:
                    if find_text in run.text:
                        run.text = run.text.replace(find_text, replace_text)
                        replaced += 1

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        for run in para.runs:
                            if find_text in run.text:
                                run.text = run.text.replace(find_text, replace_text)
                                replaced += 1

    if append_text:
        doc.add_paragraph(append_text)

    try:
        doc.save(str(p))
    except Exception as exc:
        return f"ERROR: 保存失败: {exc}"

    parts = [f"已修改 {p.name}"]
    if replaced:
        parts.append(f"替换 {replaced} 处")
    if append_text:
        parts.append(f"追加 {len(append_text)} 字符")
    return "，".join(parts)


def _docx_merge(root: Path, args: dict) -> tuple[str, bool, str]:
    files = args.get("files") or []
    out = str(args.get("out", "") or "")
    if not files or not isinstance(files, list):
        return "", False, "ERROR: merge 需要 files 列表"
    if not out:
        return "", False, "ERROR: merge 需要 out 参数"

    preview = f"将合并 {len(files)} 个文档 → {out}\n  " + "\n  ".join(files[:5])
    if len(files) > 5:
        preview += f"\n  ... 还有 {len(files) - 5} 个"
    return preview, True, ""


def _docx_merge_apply(root: Path, args: dict) -> str:
    try:
        from docx import Document
    except ImportError:
        return "ERROR: 需要 python-docx"

    files = args.get("files") or []
    out = str(args.get("out", ""))

    out_p, err = _resolve_path(root, out)
    if err:
        return err

    merged = None
    for f in files:
        fp, err = _resolve_path(root, str(f))
        if err:
            return err
        if not fp.exists():
            return f"ERROR: 文件不存在: {f}"
        try:
            doc = Document(str(fp))
        except Exception as exc:
            return f"ERROR: 打开 {f} 失败: {exc}"
        if merged is None:
            merged = doc
        else:
            for para in doc.paragraphs:
                merged.add_paragraph(para.text)
            for table in doc.tables:
                t = merged.add_table(rows=len(table.rows),
                                     cols=len(table.columns))
                for ri, row in enumerate(table.rows):
                    for ci, cell in enumerate(row.cells):
                        t.cell(ri, ci).text = cell.text

    if merged is None:
        return "ERROR: 没有可合并的文档"

    try:
        merged.save(str(out_p))
    except Exception as exc:
        return f"ERROR: 保存失败: {exc}"

    return f"已合并 {len(files)} 个文档 → {out}"


def _docx_encrypt(root: Path, args: dict) -> tuple[str, bool, str]:
    path = str(args.get("path", ""))
    password = str(args.get("password", "") or "")
    out = str(args.get("out", "") or "")
    if not password:
        return "", False, "ERROR: encrypt 需要 password"
    if not out:
        return "", False, "ERROR: encrypt 需要 out（输出路径）"

    return f"将加密 {path} → {out}", True, ""


def _docx_encrypt_apply(root: Path, args: dict) -> str:
    path = str(args.get("path", ""))
    password = str(args.get("password", ""))
    out = str(args.get("out", ""))

    src, err = _resolve_path(root, path)
    if err:
        return err
    dst, err = _resolve_path(root, out)
    if err:
        return err

    ok, msg = win32_com_available()
    if ok:
        try:
            import win32com.client
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            doc = word.Documents.Open(str(src))
            doc.SaveAs(str(dst), Password=password)
            doc.Close()
            word.Quit()
            return f"已加密（win32com）: {dst.name}"
        except Exception:
            pass

    if src.suffix.lower() == ".docx":
        return encrypt_with_msoffcrypto(src, dst, password)

    return "ERROR: 环境不支持加密（需要 Windows + Office 或 .docx 文件）"


def _docx_print(root: Path, args: dict) -> tuple[str, bool, str]:
    path = str(args.get("path", ""))
    printer = str(args.get("printer", "") or "")
    p, err = _resolve_path(root, path)
    if err:
        return "", False, err
    err = check_file(p, DOCX_EXTS)
    if err:
        return "", False, err
    preview = f"将打印 {p.name}"
    if printer:
        preview += f"  → {printer}"
    return preview, True, ""


def _docx_print_apply(root: Path, args: dict) -> str:
    path = str(args.get("path", ""))
    printer = str(args.get("printer", "") or "")
    p, err = _resolve_path(root, path)
    if err:
        return err
    return system_print(p, printer)


def _docx_info(root: Path, args: dict) -> str:
    path = str(args.get("path", ""))
    p, err = _resolve_path(root, path)
    if err:
        return err
    err = check_file(p, DOCX_EXTS)
    if err:
        return err

    stat = p.stat()
    import time
    mtime = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(stat.st_mtime))
    return (f"文件: {p.name}\n"
            f"格式: {p.suffix}\n"
            f"大小: {format_size(stat.st_size)}\n"
            f"修改时间: {mtime}")


def docx_op(root: Path, args: dict):
    op = str(args.get("op", "")).lower()

    if op == "read":
        return _docx_read(root, args)
    if op == "info":
        return _docx_info(root, args)
    if op == "extract_images":
        return _docx_extract_images(root, args)

    if op == "edit":
        return _docx_edit(root, args)
    if op == "merge":
        return _docx_merge(root, args)
    if op == "encrypt":
        return _docx_encrypt(root, args)
    if op == "print":
        return _docx_print(root, args)

    return "", False, f"ERROR: 未知 op: {op}"


def docx_apply(root: Path, args: dict) -> str:
    op = str(args.get("op", "")).lower()
    if op == "edit":
        return _docx_edit_apply(root, args)
    if op == "merge":
        return _docx_merge_apply(root, args)
    if op == "encrypt":
        return _docx_encrypt_apply(root, args)
    if op == "print":
        return _docx_print_apply(root, args)
    return f"ERROR: 未知写 op: {op}"