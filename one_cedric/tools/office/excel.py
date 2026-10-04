"""Excel 操作：openpyxl + msoffcrypto-tool。"""
from __future__ import annotations

from pathlib import Path

from ..sandbox import _resolve_path
from .common import (
    check_file, format_size, system_print,
    win32_com_available, encrypt_with_msoffcrypto,
)

XLSX_EXTS = (".xlsx", ".xlsm", ".xltx", ".xltm")


def _excel_read(root: Path, args: dict) -> str:
    try:
        import openpyxl
    except ImportError:
        return "ERROR: 需要 openpyxl（pip install openpyxl）"

    path = str(args.get("path", ""))
    sheet = str(args.get("sheet", "") or "")
    max_rows = args.get("max_rows", 50)
    max_cols = args.get("max_cols", 20)

    p, err = _resolve_path(root, path)
    if err:
        return err
    err = check_file(p, XLSX_EXTS)
    if err:
        return err

    try:
        wb = openpyxl.load_workbook(str(p), data_only=True, read_only=True)
    except Exception as exc:
        return f"ERROR: 打开失败: {exc}"

    try:
        max_r = int(max_rows) if max_rows else 50
        max_c = int(max_cols) if max_cols else 20
    except (TypeError, ValueError):
        max_r, max_c = 50, 20

    lines = [f"文件: {p.name}", f"工作表: {wb.sheetnames}", ""]

    sheets = [sheet] if sheet and sheet in wb.sheetnames else wb.sheetnames
    for sn in sheets:
        ws = wb[sn]
        lines.append(f"--- {sn} ---")
        lines.append(f"尺寸: {ws.max_row} 行 × {ws.max_column} 列")
        lines.append("")
        count = 0
        for row in ws.iter_rows(max_row=max_r, max_col=max_c, values_only=True):
            cells = []
            for v in row:
                if v is None:
                    cells.append("")
                else:
                    s = str(v)
                    cells.append(s if len(s) <= 30 else s[:27] + "…")
            lines.append(" | ".join(cells))
            count += 1
        if ws.max_row > max_r:
            lines.append(f"... 还有 {ws.max_row - max_r} 行")
        lines.append("")

    return "\n".join(lines)


def _excel_info(root: Path, args: dict) -> str:
    try:
        import openpyxl
    except ImportError:
        return "ERROR: 需要 openpyxl"

    path = str(args.get("path", ""))
    p, err = _resolve_path(root, path)
    if err:
        return err
    err = check_file(p, XLSX_EXTS)
    if err:
        return err

    try:
        wb = openpyxl.load_workbook(str(p), read_only=True)
    except Exception as exc:
        return f"ERROR: 打开失败: {exc}"

    stat = p.stat()
    import time
    mtime = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(stat.st_mtime))

    lines = [f"文件: {p.name}",
             f"大小: {format_size(stat.st_size)}",
             f"修改时间: {mtime}",
             f"工作表数: {len(wb.sheetnames)}",
             "工作表:"]
    for sn in wb.sheetnames:
        ws = wb[sn]
        lines.append(f"  {sn}: {ws.max_row} 行 × {ws.max_column} 列")
    return "\n".join(lines)


def _excel_worksheets(root: Path, args: dict) -> str:
    return _excel_info(root, args)


def _excel_list_sheets(root: Path, args: dict) -> str:
    try:
        import openpyxl
    except ImportError:
        return "ERROR: 需要 openpyxl"

    path = str(args.get("path", ""))
    p, err = _resolve_path(root, path)
    if err:
        return err
    err = check_file(p, XLSX_EXTS)
    if err:
        return err

    wb = openpyxl.load_workbook(str(p), read_only=True)
    return "工作表:\n" + "\n".join(f"  {i+1}. {n}" for i, n in enumerate(wb.sheetnames))


def _excel_formulas(root: Path, args: dict) -> str:
    try:
        import openpyxl
    except ImportError:
        return "ERROR: 需要 openpyxl"

    path = str(args.get("path", ""))
    sheet = str(args.get("sheet", "") or "")
    p, err = _resolve_path(root, path)
    if err:
        return err
    err = check_file(p, XLSX_EXTS)
    if err:
        return err

    try:
        wb = openpyxl.load_workbook(str(p), data_only=False)
    except Exception as exc:
        return f"ERROR: 打开失败: {exc}"

    sheets = [sheet] if sheet and sheet in wb.sheetnames else wb.sheetnames
    lines = [f"文件: {p.name} 的公式", ""]
    for sn in sheets:
        ws = wb[sn]
        lines.append(f"--- {sn} ---")
        count = 0
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    lines.append(f"  {cell.coordinate}: {cell.value}")
                    count += 1
                    if count >= 200:
                        break
            if count >= 200:
                break
        lines.append(f"  （{count} 个公式）")
        lines.append("")
    return "\n".join(lines)


def _excel_edit(root: Path, args: dict):
    path = str(args.get("path", ""))
    sheet = str(args.get("sheet", "") or "")
    cell = str(args.get("cell", "") or "")
    value = args.get("value")
    if not cell:
        return "", False, "ERROR: edit 需要 cell（如 'A1'）"
    return f"将在 {path} 的 {sheet or '当前表'} 的 {cell} 写入 {value!r}", True, ""


def _excel_edit_apply(root: Path, args: dict) -> str:
    try:
        import openpyxl
    except ImportError:
        return "ERROR: 需要 openpyxl"

    path = str(args.get("path", ""))
    sheet = str(args.get("sheet", "") or "")
    cell = str(args.get("cell", ""))
    value = args.get("value")

    p, err = _resolve_path(root, path)
    if err:
        return err

    try:
        wb = openpyxl.load_workbook(str(p))
    except Exception as exc:
        return f"ERROR: 打开失败: {exc}"

    if not sheet:
        sheet = wb.sheetnames[0]
    if sheet not in wb.sheetnames:
        return f"ERROR: 工作表不存在: {sheet}"

    ws = wb[sheet]
    try:
        ws[cell] = value
    except Exception as exc:
        return f"ERROR: 写入失败: {exc}"

    try:
        wb.save(str(p))
    except Exception as exc:
        return f"ERROR: 保存失败: {exc}"

    return f"已写入 {sheet}!{cell} = {value!r}"


def _excel_merge(root: Path, args: dict):
    files = args.get("files") or []
    out = str(args.get("out", "") or "")
    if not files or not isinstance(files, list):
        return "", False, "ERROR: merge 需要 files 列表"
    if not out:
        return "", False, "ERROR: merge 需要 out"
    return f"将合并 {len(files)} 个工作簿 → {out}", True, ""


def _excel_merge_apply(root: Path, args: dict) -> str:
    try:
        import openpyxl
    except ImportError:
        return "ERROR: 需要 openpyxl"

    files = args.get("files") or []
    out = str(args.get("out", ""))

    out_p, err = _resolve_path(root, out)
    if err:
        return err

    merged = openpyxl.Workbook()
    merged.remove(merged.active)

    for f in files:
        fp, err = _resolve_path(root, str(f))
        if err:
            return err
        try:
            src = openpyxl.load_workbook(str(fp), data_only=False)
        except Exception as exc:
            return f"ERROR: 打开 {f} 失败: {exc}"
        for sn in src.sheetnames:
            ws_src = src[sn]
            ws_dst = merged.create_sheet(title=f"{fp.stem}_{sn}"[:31])
            for row in ws_src.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        ws_dst[cell.coordinate] = cell.value

    try:
        merged.save(str(out_p))
    except Exception as exc:
        return f"ERROR: 保存失败: {exc}"
    return f"已合并 {len(files)} 个工作簿 → {out}"


def _excel_encrypt(root: Path, args: dict):
    path = str(args.get("path", ""))
    password = str(args.get("password", "") or "")
    out = str(args.get("out", "") or "")
    if not password:
        return "", False, "ERROR: encrypt 需要 password"
    if not out:
        return "", False, "ERROR: encrypt 需要 out"
    return f"将加密 {path} → {out}", True, ""


def _excel_encrypt_apply(root: Path, args: dict) -> str:
    path = str(args.get("path", ""))
    password = str(args.get("password", ""))
    out = str(args.get("out", ""))

    src, err = _resolve_path(root, path)
    if err:
        return err
    dst, err = _resolve_path(root, out)
    if err:
        return err

    ok, _ = win32_com_available()
    if ok:
        try:
            import win32com.client
            excel = win32com.client.Dispatch("Excel.Application")
            excel.Visible = False
            wb = excel.Workbooks.Open(str(src))
            wb.SaveAs(str(dst), Password=password)
            wb.Close()
            excel.Quit()
            return f"已加密（win32com）: {dst.name}"
        except Exception:
            pass

    return encrypt_with_msoffcrypto(src, dst, password)


def _excel_print(root: Path, args: dict):
    path = str(args.get("path", ""))
    printer = str(args.get("printer", "") or "")
    p, err = _resolve_path(root, path)
    if err:
        return "", False, err
    err = check_file(p, XLSX_EXTS)
    if err:
        return "", False, err
    preview = f"将打印 {p.name}"
    if printer:
        preview += f"  → {printer}"
    return preview, True, ""


def _excel_print_apply(root: Path, args: dict) -> str:
    path = str(args.get("path", ""))
    printer = str(args.get("printer", "") or "")
    p, err = _resolve_path(root, path)
    if err:
        return err
    return system_print(p, printer)


def excel_op(root: Path, args: dict):
    op = str(args.get("op", "")).lower()
    if op == "read":
        return _excel_read(root, args)
    if op == "info":
        return _excel_info(root, args)
    if op == "worksheets":
        return _excel_worksheets(root, args)
    if op == "list_sheets":
        return _excel_list_sheets(root, args)
    if op == "formulas":
        return _excel_formulas(root, args)

    if op == "edit":
        return _excel_edit(root, args)
    if op == "merge":
        return _excel_merge(root, args)
    if op == "encrypt":
        return _excel_encrypt(root, args)
    if op == "print":
        return _excel_print(root, args)

    return "", False, f"ERROR: 未知 op: {op}"


def excel_apply(root: Path, args: dict) -> str:
    op = str(args.get("op", "")).lower()
    if op == "edit":
        return _excel_edit_apply(root, args)
    if op == "merge":
        return _excel_merge_apply(root, args)
    if op == "encrypt":
        return _excel_encrypt_apply(root, args)
    if op == "print":
        return _excel_print_apply(root, args)
    return f"ERROR: 未知写 op: {op}"