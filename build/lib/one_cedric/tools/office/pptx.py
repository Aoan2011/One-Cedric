"""PowerPoint 操作：python-pptx。"""
from __future__ import annotations

from pathlib import Path

from ..sandbox import _resolve_path
from .common import check_file, format_size, system_print, win32_com_available


def _pptx_read(root: Path, args: dict) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        return "ERROR: 需要 python-pptx（pip install python-pptx）"

    path = str(args.get("path", ""))
    p, err = _resolve_path(root, path)
    if err:
        return err
    err = check_file(p, (".pptx", ".ppt"))
    if err:
        return err

    try:
        prs = Presentation(str(p))
    except Exception as exc:
        return f"ERROR: 打开失败: {exc}"

    lines = [f"文件: {p.name}", f"幻灯片数: {len(prs.slides)}", ""]
    for si, slide in enumerate(prs.slides, 1):
        lines.append(f"--- 第 {si} 页 ---")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    lines.append(text[:500])
            elif shape.shape_type == 13:
                lines.append(f"[图片] {shape.name}")
        lines.append("")
        if si >= 30:
            lines.append(f"... 还有 {len(prs.slides) - 30} 页")
            break
    return "\n".join(lines)


def _pptx_info(root: Path, args: dict) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        return "ERROR: 需要 python-pptx"

    path = str(args.get("path", ""))
    p, err = _resolve_path(root, path)
    if err:
        return err
    err = check_file(p, (".pptx", ".ppt"))
    if err:
        return err

    try:
        prs = Presentation(str(p))
    except Exception as exc:
        return f"ERROR: 打开失败: {exc}"

    stat = p.stat()
    import time
    mtime = time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(stat.st_mtime))
    size = prs.slide_width
    height = prs.slide_height
    return (f"文件: {p.name}\n"
            f"幻灯片数: {len(prs.slides)}\n"
            f"尺寸: {size} × {height} EMU\n"
            f"文件大小: {format_size(stat.st_size)}\n"
            f"修改时间: {mtime}")


def _pptx_extract_content(root: Path, args: dict) -> str:
    return _pptx_read(root, args)


def _pptx_edit(root: Path, args: dict):
    path = str(args.get("path", ""))
    find_text = str(args.get("find", "") or "")
    replace_text = str(args.get("replace", "") or "")
    if not find_text:
        return "", False, "ERROR: edit 需要 find 参数"
    return f"将在 {path} 中替换 '{find_text[:50]}' → '{replace_text[:50]}'", True, ""


def _pptx_edit_apply(root: Path, args: dict) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        return "ERROR: 需要 python-pptx"

    path = str(args.get("path", ""))
    find_text = str(args.get("find", ""))
    replace_text = str(args.get("replace", ""))

    p, err = _resolve_path(root, path)
    if err:
        return err

    try:
        prs = Presentation(str(p))
    except Exception as exc:
        return f"ERROR: 打开失败: {exc}"

    replaced = 0
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if find_text in run.text:
                        run.text = run.text.replace(find_text, replace_text)
                        replaced += 1

    try:
        prs.save(str(p))
    except Exception as exc:
        return f"ERROR: 保存失败: {exc}"

    return f"已替换 {replaced} 处"


def _pptx_merge(root: Path, args: dict):
    files = args.get("files") or []
    out = str(args.get("out", "") or "")
    if not files or not isinstance(files, list):
        return "", False, "ERROR: merge 需要 files 列表"
    if not out:
        return "", False, "ERROR: merge 需要 out"
    preview = f"将合并 {len(files)} 个 PPT → {out}"
    return preview, True, ""


def _pptx_merge_apply(root: Path, args: dict) -> str:
    try:
        from pptx import Presentation
    except ImportError:
        return "ERROR: 需要 python-pptx"

    files = args.get("files") or []
    out = str(args.get("out", ""))

    out_p, err = _resolve_path(root, out)
    if err:
        return err

    if not files:
        return "ERROR: files 为空"

    first_p, err = _resolve_path(root, str(files[0]))
    if err:
        return err
    try:
        merged = Presentation(str(first_p))
    except Exception as exc:
        return f"ERROR: 打开 {files[0]} 失败: {exc}"

    for f in files[1:]:
        fp, err = _resolve_path(root, str(f))
        if err:
            return err
        try:
            src = Presentation(str(fp))
        except Exception as exc:
            return f"ERROR: 打开 {f} 失败: {exc}"
        for slide in src.slides:
            blank = merged.slide_layouts[6] if len(merged.slide_layouts) > 6 else merged.slide_layouts[0]
            new_slide = merged.slides.add_slide(blank)
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    tb = new_slide.shapes.add_textbox(0, 0, merged.slide_width, merged.slide_height)
                    tb.text_frame.text = shape.text_frame.text

    try:
        merged.save(str(out_p))
    except Exception as exc:
        return f"ERROR: 保存失败: {exc}"

    return f"已合并 {len(files)} 个 PPT → {out}"


def _pptx_encrypt(root: Path, args: dict):
    path = str(args.get("path", ""))
    password = str(args.get("password", "") or "")
    out = str(args.get("out", "") or "")
    if not password:
        return "", False, "ERROR: encrypt 需要 password"
    if not out:
        return "", False, "ERROR: encrypt 需要 out"
    return f"将加密 {path} → {out}", True, ""


def _pptx_encrypt_apply(root: Path, args: dict) -> str:
    path = str(args.get("path", ""))
    password = str(args.get("password", ""))
    out = str(args.get("out", ""))

    src, err = _resolve_path(root, path)
    if err:
        return err
    dst, err = _resolve_path(root, out)
    if err:
        return err

    ok, _ = win32_com_available()
    if not ok:
        return "ERROR: 加密需要 Windows + Office（win32com）"

    try:
        import win32com.client
        ppt = win32com.client.Dispatch("PowerPoint.Application")
        pres = ppt.Presentations.Open(str(src), WithWindow=False)
        pres.SaveAs(str(dst), Password=password)
        pres.Close()
        ppt.Quit()
        return f"已加密: {dst.name}"
    except Exception as exc:
        return f"ERROR: 加密失败: {exc}"


def _pptx_print(root: Path, args: dict):
    path = str(args.get("path", ""))
    printer = str(args.get("printer", "") or "")
    p, err = _resolve_path(root, path)
    if err:
        return "", False, err
    err = check_file(p, (".pptx", ".ppt"))
    if err:
        return "", False, err
    preview = f"将打印 {p.name}"
    if printer:
        preview += f"  → {printer}"
    return preview, True, ""


def _pptx_print_apply(root: Path, args: dict) -> str:
    path = str(args.get("path", ""))
    printer = str(args.get("printer", "") or "")
    p, err = _resolve_path(root, path)
    if err:
        return err
    return system_print(p, printer)


def pptx_op(root: Path, args: dict):
    op = str(args.get("op", "")).lower()
    if op == "read":
        return _pptx_read(root, args)
    if op == "info":
        return _pptx_info(root, args)
    if op == "extract_content":
        return _pptx_extract_content(root, args)

    if op == "edit":
        return _pptx_edit(root, args)
    if op == "merge":
        return _pptx_merge(root, args)
    if op == "encrypt":
        return _pptx_encrypt(root, args)
    if op == "print":
        return _pptx_print(root, args)

    return "", False, f"ERROR: 未知 op: {op}"


def pptx_apply(root: Path, args: dict) -> str:
    op = str(args.get("op", "")).lower()
    if op == "edit":
        return _pptx_edit_apply(root, args)
    if op == "merge":
        return _pptx_merge_apply(root, args)
    if op == "encrypt":
        return _pptx_encrypt_apply(root, args)
    if op == "print":
        return _pptx_print_apply(root, args)
    return f"ERROR: 未知写 op: {op}"